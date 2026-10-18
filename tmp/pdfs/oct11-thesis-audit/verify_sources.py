from pathlib import Path
import ast
import collections
import hashlib
import json
import math
import re
import statistics as st

import openpyxl
import pypdfium2 as pdfium
from pypdf import PdfReader
from PIL import Image, ImageDraw

ROOT = Path(r"C:\Users\justine mendoza\source\repos\NFC_System")
OUT = Path(__file__).resolve().parent
DATA = ROOT / "output/chapter4"
DOCS = Path(r"C:\Users\justine mendoza\Documents\METRICS first batch")
CATS = [("Functional Suitability", 0, 5), ("Performance Efficiency", 5, 9),
        ("Usability", 9, 14), ("Reliability", 14, 19), ("Security", 19, 24)]
result = {"sources": [], "surveys": {}, "metrics": {}}

def track(path):
    result["sources"].append({"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    return path

def describe(values):
    return dict(n=len(values), mean=st.mean(values), median=st.median(values),
                sample_sd=st.stdev(values), minimum=min(values), maximum=max(values))

source = track(DATA / "Survey-Response-Entry-Template.xlsx")
book = openpyxl.load_workbook(source, read_only=True, data_only=False)
answers = {}
for group in ["Manual", "System"]:
    rows = [list(row[1:25]) for row in book[group].iter_rows(min_row=2, max_row=71, values_only=True)]
    assert len(rows) == 70 and all(len(row) == 24 and all(v in [1, 2, 3, 4] for v in row) for row in rows)
    answers[group] = rows
    scores = {name: [st.mean(row[a:b]) for row in rows] for name, a, b in CATS}
    scores["Overall"] = [st.mean(scores[name][i] for name, _, _ in CATS) for i in range(70)]
    result["surveys"][group] = {name: describe(values) for name, values in scores.items()}
book.close()
manual, system = [result["surveys"][g]["Overall"] for g in ["Manual", "System"]]
a, b = manual["sample_sd"] ** 2 / 70, system["sample_sd"] ** 2 / 70
diff = system["mean"] - manual["mean"]
se = math.sqrt(a + b)
t = diff / se
df = (a + b) ** 2 / ((a * a + b * b) / 69)
tree = ast.parse((DATA / "analyze_completed_survey.py").read_text())
scope = {"math": math}
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in
                            ["beta_fraction", "regularized_beta", "two_sided_p", "critical_t"]], type_ignores=[]),
             "distribution_helpers", "exec"), scope)
critical = scope["critical_t"](df)
result["welch"] = dict(difference=diff, se=se, t=t, df=df, p=scope["two_sided_p"](t, df),
                       ci95=[diff-critical*se, diff+critical*se])

survey_pdf = track(DOCS / "UnStudent Weighted Mean - Students (1).pdf")
reader = PdfReader(survey_pdf)
left = "\n".join(p.extract_text() for p in reader.pages[:9])
middle = "\n".join(p.extract_text() for p in reader.pages[9:18])
questions = list(re.finditer(r'(?<!\d)([1-5])\.\s+(.*?)(\d+)\s+(\d+)\s*(?=\n|$)', left, re.S))
numeric = list(re.finditer(r'(?m)^[ \t]*(\d+)[ \t]+(\d+)[ \t]+(\d+)[ \t]+(\d+)[ \t]+([\d.]+)(?:[ \t]+([\d.]+))?[ \t]*$', middle))
assert len(questions) == len(numeric) == 111
personnel = []
for start, count, label, lengths in [(24, 24, "Security Personnel", [5,4,5,5,5]),
                                     (48, 39, "Faculty/Event Organizer/Admin", [5,5,4,5,5,5,5,5])]:
    items = []
    for i in range(start, start+count):
        f4, f3 = int(questions[i][3]), int(questions[i][4])
        f2, f1, n, total = map(int, numeric[i].groups()[:4])
        values = [rating for rating, freq in zip([4,3,2,1], [f4,f3,f2,f1]) for _ in range(freq)]
        assert len(values) == n == 5 and sum(values) == total
        items.append(describe(values))
    categories = []
    offset = 0
    for length in lengths:
        current = items[offset:offset+length]
        categories.append(dict(mean=st.mean(i["mean"] for i in current),
                               sd_range=[min(i["sample_sd"] for i in current), max(i["sample_sd"] for i in current)]))
        offset += length
    personnel.append(dict(group=label, categories=categories,
                          overall_equal_category_mean=st.mean(c["mean"] for c in categories), items=items))
result["personnel_surveys"] = personnel

records, runs = [], []
for batch, filename in enumerate(["METRICS FIRST BATCH.xlsx", "Second Batch.xlsx"], 1):
    w = openpyxl.load_workbook(track(DOCS / filename), read_only=True, data_only=True)
    for sheet in w:
        rows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] is not None]
        records.extend(rows)
        segments = []
        for r in sorted(rows, key=lambda r: r[0]):
            if r[4] != "GRANTED":
                continue
            if not segments or segments[-1][0][2] != r[2]:
                segments.append([])
            segments[-1].append(r)
        assert [s[0][2] for s in segments] == ["Entry", "Exit", "Entry", "Exit"]
        assert [len(s) for s in segments] == ([20]*4 if batch == 1 else [20,20,10,10])
        for s in [s for s in segments if s[0][2] == "Entry"]:
            runs.append(dict(mode=s[0][3], n=len(s), intervals=len(s)-1,
                             span=(s[-1][0]-s[0][0]).total_seconds(),
                             denied_within=sum(r[4] != "GRANTED" and s[0][0] <= r[0] <= s[-1][0] for r in rows)))
    w.close()
assert len(records) == 438 and len(set(records)) == 438
for mode in ["Fast", "Standard", "HighSecurity"]:
    selected = [r for r in records if r[3] == mode]
    entries = [r for r in selected if r[2] == "Entry" and r[4] == "GRANTED"]
    assert len(entries) == 70
    assert all(abs(r[11]-sum(r[i] if isinstance(r[i],(int,float)) else 0 for i in [5,7,9])) < 1e-8 for r in entries)
    rr = [r for r in runs if r["mode"] == mode]
    span = sum(r["span"] for r in rr)
    intervals = sum(r["intervals"] for r in rr)
    result["metrics"][mode] = dict(records=len(selected), entries=len(entries),
        exits=sum(r[2] == "Exit" and r[4] == "GRANTED" for r in selected),
        denials=sum(r[4] != "GRANTED" for r in selected), internal_ms=describe([r[11] for r in entries]),
        db_ms=describe([r[12] for r in entries]), intervals=intervals, span_s=span,
        spacing_s=span/intervals, interval_rate_per_min=60*intervals/span,
        denied_within=sum(r["denied_within"] for r in rr))
    if mode != "Fast":
        result["metrics"][mode]["interaction_s"] = describe([r[10]/1000 for r in entries])
        result["metrics"][mode]["pin_s"] = describe([r[6]/1000 for r in entries])
        if mode == "HighSecurity":
            result["metrics"][mode]["qr_s"] = describe([r[8]/1000 for r in entries])
result["denial_codes"] = dict(collections.Counter(r[4] for r in records if r[4] != "GRANTED"))
applicable = [r for r in records if (r[2] == "Entry" and r[4] == "GRANTED") or
              r[4] in ["DENIED [PIN_FAILURE]", "DENIED [PIN_LOCKED]", "DENIED [QR_REPLACED_OR_MISMATCH]"]]
assert len(applicable) == 222
pin = [r for r in applicable if r[3] != "Fast"]
assert len(pin) == 152
result["security_denominators"] = dict(applicable_credential_records=222, recorded_pin_submissions=152,
    pin_locked_records=sum(r[4] == "DENIED [PIN_LOCKED]" for r in pin),
    caveat="Expected outcome labels and newly-triggered status depend on researcher confirmation, not exports alone.")

manual_pdf = track(DOCS / "Manual Process Baseline Testing - Sheet1.pdf")
text = "\n".join(p.extract_text() for p in PdfReader(manual_pdf).pages)
values = [float(v) for _,v in re.findall(r"Respondent\s+(\d+)\s+([\d.]+)", text)]
assert len(values) == 70
result["manual_writing_s"] = describe(values)
result["manual_writing_s"].update(total=sum(values), service_equivalent_per_min=60/st.mean(values))
(OUT / "source-verification.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
print(json.dumps({k:v for k,v in result.items() if k not in ["sources", "personnel_surveys"]}, indent=2))
print("Personnel overall means:", [(g["group"],g["overall_equal_category_mean"]) for g in personnel])

pdf = pdfium.PdfDocument(r"C:\Users\justine mendoza\Documents\NFC Thesis Chapter 1-5 (1).pdf")
pages = [21,81,85,88,98,106,109,66,67,161,117,119,125,159,214,230,231,234,235,236,237,218,219,220,221]
for start in range(0,len(pages),4):
    contact = Image.new("RGB", (1400,1950), "#dedede")
    draw = ImageDraw.Draw(contact)
    for j, number in enumerate(pages[start:start+4]):
        picture = pdf[number-1].render(scale=1.5).to_pil()
        picture.save(OUT / f"page-{number}.png")
        picture.thumbnail((680,920))
        x,y = 10+(j%2)*700, 35+(j//2)*975
        contact.paste(picture,(x,y))
        draw.text((x,y-25),f"Page {number}",fill="black")
    contact.save(OUT / f"contact-{start//4+1}.png")
print("Rendered", len(pages), "selected pages.")
