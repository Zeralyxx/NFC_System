import json
import re
import statistics as st
from pathlib import Path
import openpyxl
import pdfplumber

ROOT = Path(r"C:\Users\justine mendoza\source\repos\NFC_System")
OUT = Path(__file__).resolve().parent
verified = json.loads((OUT / "source-verification.json").read_text())
expected = []
book = openpyxl.load_workbook(ROOT / "output/chapter4/Survey-Response-Entry-Template.xlsx", read_only=True, data_only=False)
for name in ["System"]:
    rows = [list(r[1:25]) for r in book[name].iter_rows(min_row=2,max_row=71,values_only=True)]
    expected.extend((st.mean(v),st.stdev(v)) for v in zip(*rows))
for group in verified["personnel_surveys"]:
    expected.extend((v["mean"],v["sample_sd"]) for v in group["items"])
rows = [list(r[1:25]) for r in book["Manual"].iter_rows(min_row=2,max_row=71,values_only=True)]
expected.extend((st.mean(v),st.stdev(v)) for v in zip(*rows))
book.close()
found = []
with pdfplumber.open(r"C:\Users\justine mendoza\Documents\NFC Thesis Chapter 1-5 (1).pdf") as pdf:
    for number in range(135,159):
        for table in pdf.pages[number-1].extract_tables():
            for row in table:
                if len(row) == 4 and re.match(r"^[1-5]\.\s",row[0] or ""):
                    try:
                        found.append((number,float(row[1]),float(row[2]),row[0]))
                    except (TypeError,ValueError):
                        pass
assert len(found) == len(expected) == 111, (len(found),len(expected))
errors = []
for i, ((page,mean,sd,question),(a,b)) in enumerate(zip(found,expected),1):
    if abs(mean-a)>0.00501 or abs(sd-b)>0.00501:
        errors.append(dict(item=i,page=page,question=question,pdf_mean=mean,pdf_sd=sd,expected_mean=a,expected_sd=b))
report = dict(checked=111,errors=errors)
(OUT / "item-table-check.json").write_text(json.dumps(report,indent=2),encoding="utf-8")
print(json.dumps(report,indent=2))
