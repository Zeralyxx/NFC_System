import json
import math
import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from lxml import etree

folder = Path(__file__).resolve().parent
original = Path(r'C:\Users\justine mendoza\Documents\METRICS first batch\UnStudent-Weighted-Mean.xlsx')
system_rows = [r for a, b in [(5, 9), (13, 16), (20, 24), (28, 32), (36, 40)] for r in range(a, b + 1)]
manual_rows = [r for a, b in [(168, 172), (176, 179), (183, 187), (191, 195), (199, 203)] for r in range(a, b + 1)]
category_rows = [10, 17, 25, 33, 41, 43, 173, 180, 188, 196, 204, 206]
if '--prepare' in sys.argv:
    wb = openpyxl.load_workbook(original, data_only=False)
    addresses = [f'{c}{r}' for r in system_rows + manual_rows for c in 'FGHIJ']
    addresses += [f'{c}{r}' for r in category_rows for c in 'HIJ']
    formulas = {address: wb['Students'][address].value for address in addresses}
    assert all(isinstance(value, str) and value.startswith('=') for value in formulas.values())
    (folder / 'tally-original-formulas.json').write_text(json.dumps(formulas), encoding='utf-8')
    print('Captured original formulas, including shared-formula children.')
    sys.exit(0)

manifest = json.loads((folder / 'tally-update-manifest.json').read_text(encoding='utf-8'))
working = openpyxl.load_workbook(folder / 'tally-update-working.xlsx', data_only=True)
source = openpyxl.load_workbook(folder / 'Survey-Response-Entry-Template.xlsx', data_only=False)
before = openpyxl.load_workbook(original, data_only=False)
sheet = working['Students']
for name, target_rows in [('System', system_rows), ('Manual', manual_rows)]:
    raw = [[source[name].cell(r, c).value for c in range(2, 26)] for r in range(2, 72)]
    for i, r in enumerate(target_rows):
        counts = [sum(row[i] == rating for row in raw) for rating in (4, 3, 2, 1)]
        assert [sheet.cell(r, c).value for c in range(2, 6)] == counts
        assert sheet[f'F{r}'].value == 70
        total = sum(count * rating for count, rating in zip(counts, (4, 3, 2, 1)))
        assert sheet[f'G{r}'].value == total, (r, total, sheet[f'G{r}'].value)
        # Excel rounds half away from zero; all means are positive.
        rounded = math.floor(total / 70 * 100 + 0.5) / 100
        assert sheet[f'H{r}'].value == rounded, (r, rounded, sheet[f'H{r}'].value)
assert sheet['H43'].value == 3.74
assert sheet['H206'].value == 2.28

ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
tag = '{' + ns['s'] + '}'
updated_addresses = set(manifest['inputCells'] + manifest['dependentCells'])
backup = folder / 'UnStudent-Weighted-Mean-before-tally-update.xlsx'
if not backup.exists():
    shutil.copy2(original, backup)
output = folder / 'UnStudent-Weighted-Mean-Corrected.xlsx'
with ZipFile(original) as zin, ZipFile(output, 'w') as zout:
    for info in zin.infolist():
        content = zin.read(info.filename)
        if info.filename == 'xl/worksheets/sheet1.xml':
            root = etree.fromstring(content)
            cells = {cell.get('r'): cell for cell in root.findall('.//s:sheetData/s:row/s:c', ns)}
            for address in updated_addresses:
                cell = cells[address]
                value = sheet[address].value
                assert value is not None and not (isinstance(value, str) and value.startswith('#')), (address, value)
                for child in list(cell):
                    if child.tag in (tag + 'v', tag + 'is'):
                        cell.remove(child)
                cell.attrib.pop('t', None)
                cache = etree.SubElement(cell, tag + 'v')
                if isinstance(value, str):
                    cell.set('t', 'str')
                    cache.text = value
                else:
                    cache.text = str(value)
            content = etree.tostring(root, encoding='utf-8')
        zout.writestr(info, content)

after = openpyxl.load_workbook(output, data_only=False)
assert before.sheetnames == after.sheetnames
for old_sheet, new_sheet in zip(before, after):
    assert list(map(str, old_sheet.merged_cells.ranges)) == list(map(str, new_sheet.merged_cells.ranges))
    assert old_sheet.freeze_panes == new_sheet.freeze_panes
    for row in old_sheet:
        for cell in row:
            target = new_sheet[cell.coordinate]
            if old_sheet.title == 'Students' and cell.coordinate in manifest['inputCells']:
                assert target.value == sheet[cell.coordinate].value
            else:
                assert cell.value == target.value, (old_sheet.title, cell.coordinate, cell.value, target.value)
            assert cell._style == target._style, (old_sheet.title, cell.coordinate, 'style')
with ZipFile(original) as a, ZipFile(output) as b:
    assert a.namelist() == b.namelist()
    assert all(a.read(name) == b.read(name) for name in a.namelist() if name != 'xl/worksheets/sheet1.xml')
print(f'Verified 48 tallies, each with N=70; means 3.74 (system) and 2.28 (manual).')
print('Original formulas, styles, sheet order, and all unrelated package parts preserved.')
print(output)
