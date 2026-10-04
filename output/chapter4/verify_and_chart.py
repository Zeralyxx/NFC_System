import datetime as dt
import json
import math
from pathlib import Path

import numpy as np
import openpyxl
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).parent
data = json.loads((ROOT / 'analysis.json').read_text())
book = openpyxl.load_workbook(ROOT / 'Chapter-4-System-Metrics.xlsx', data_only=True)
errors = [(s.title, c.coordinate, c.value) for s in book for row in s for c in row if c.data_type == 'e']
assert not errors, errors
for i, record in enumerate(data['records'], 7):
    actual = [book['RawRecords'].cell(i, j).value for j in range(1, 16)]
    expected = record['values']
    assert actual[0] == dt.datetime.fromisoformat(expected[0]), (i, actual[0], expected[0])
    for j in range(1, 15):
        if isinstance(expected[j], (int, float)):
            assert math.isclose(actual[j], expected[j], abs_tol=1e-9)
        else:
            assert actual[j] == expected[j]
for i, perf in enumerate(data['performance']):
    assert book['Summary'].cell(7 + i, 2).value == 70
    assert math.isclose(book['Summary'].cell(15 + i, 3).value, perf['metrics']['total_system_ms']['mean'], abs_tol=1e-9)
    assert math.isclose(book['Summary'].cell(15 + i, 5).value, perf['metrics']['total_system_ms']['sd'], abs_tol=1e-9)
    assert math.isclose(book['Summary'].cell(29 + i, 6).value, data['pooled_spacing'][i]['rate_per_min'], abs_tol=1e-6)
assert book['Summary']['E10'].value == 438
book.close()
print('Verified: all 438 original rows, cached formulas, timing statistics and recorded-entry rates; no Excel error cells.')

FONT = Path('C:/Windows/Fonts')
def font(size, bold=False):
    return ImageFont.truetype(str(FONT / ('arialbd.ttf' if bold else 'arial.ttf')), size)
INK, GRAY = '#20282A', '#5B666B'
COLORS = ['#178072', '#4274B7', '#AC5061']

def text(draw, xy, value, size=28, color=INK, bold=False, anchor=None):
    draw.text(xy, value, font=font(size, bold), fill=color, anchor=anchor)

def axes(draw, left, top, right, bottom, ymax, ticks, label):
    for tick in ticks:
        y = bottom - tick / ymax * (bottom - top)
        draw.line((left, y, right, y), fill='#E1E5E7', width=2)
        text(draw, (left - 18, y), str(tick), 25, GRAY, anchor='rm')
    draw.line((left, top, left, bottom, right, bottom), fill=INK, width=3)
    text(draw, (left, top - 42), label, 26, GRAY)
    return lambda value: bottom - value / ymax * (bottom - top)

im = Image.new('RGB', (2200, 1100), 'white')
draw = ImageDraw.Draw(im)
text(draw, (90, 50), 'Successful-entry timing by verification mode', 46, bold=True)
text(draw, (90, 115), '70 successful entries per mode | 12-13 August 2026', 29, GRAY)
text(draw, (100, 210), 'A. Mean internal system latency', 34, bold=True)
text(draw, (1170, 210), 'B. Recorded PIN/QR interaction duration', 34, bold=True)
ys = axes(draw, 120, 330, 1010, 850, 90, range(0, 91, 15), 'Milliseconds')
labels = ['Fast', 'Standard', 'High-Security']
for i, perf in enumerate(data['performance']):
    x = 270 + i * 295
    mean, sd = [perf['metrics']['total_system_ms'][k] for k in ['mean', 'sd']]
    draw.rectangle((x - 72, ys(mean), x + 72, ys(0)), fill=COLORS[i])
    draw.line((x, ys(mean + sd), x, ys(max(0, mean - sd))), fill=INK, width=4)
    for v in [mean + sd, max(0, mean - sd)]:
        draw.line((x - 24, ys(v), x + 24, ys(v)), fill=INK, width=4)
    text(draw, (x, ys(mean + sd) - 25), f'{mean:.2f}', 30, bold=True, anchor='mb')
    text(draw, (x, 890), labels[i], 29, anchor='mt')
text(draw, (120, 964), 'Error bars show sample SD, not confidence intervals.', 26, GRAY)

yu = axes(draw, 1210, 330, 2080, 850, 24, range(0, 25, 4), 'Seconds')
for i, mode in enumerate(['Standard', 'HighSecurity']):
    values = [r['values'][10] / 1000 for r in data['records'] if r['category'] == 'Successful entry' and r['values'][3] == mode]
    x = 1450 + i * 420
    q1, median, q3 = np.percentile(values, [25, 50, 75])
    low = min(v for v in values if v >= q1 - 1.5 * (q3 - q1))
    high = max(v for v in values if v <= q3 + 1.5 * (q3 - q1))
    draw.line((x, yu(high), x, yu(low)), fill=GRAY, width=3)
    for v in [low, high]: draw.line((x - 30, yu(v), x + 30, yu(v)), fill=GRAY, width=3)
    draw.rectangle((x - 74, yu(q3), x + 74, yu(q1)), fill=COLORS[i + 1])
    draw.line((x - 74, yu(median), x + 74, yu(median)), fill='white', width=5)
    for j, value in enumerate(values):
        dx = ((j * 37) % 59 - 29) * 2
        y = yu(value)
        draw.ellipse((x + dx - 4, y - 4, x + dx + 4, y + 4), fill=INK)
    text(draw, (x, 890), labels[i + 1], 29, anchor='mt')
    text(draw, (x, 940), f'Mean {np.mean(values):.3f} s | Median {median:.3f} s', 25, GRAY, anchor='mt')
text(draw, (1210, 1000), 'Fast bypasses PIN/QR; no user duration was logged.', 26, GRAY)
text(draw, (90, 1060), 'Internal latency and interaction timers do not measure the complete physical entry process.', 27, GRAY)
im.save(ROOT / 'Figure-System-Timing.png')

im = Image.new('RGB', (1650, 1000), 'white')
draw = ImageDraw.Draw(im)
text(draw, (85, 50), 'Recorded successful-entry rate', 46, bold=True)
text(draw, (85, 115), 'Derived from first-to-last log timestamps within four entry runs per mode', 27, GRAY)
yr = axes(draw, 150, 255, 1500, 755, 25, range(0, 26, 5), 'Recorded entry intervals per minute')
for i, item in enumerate(data['pooled_spacing']):
    x = 370 + 460 * i
    draw.rectangle((x - 105, yr(item['rate_per_min']), x + 105, yr(0)), fill=COLORS[i])
    text(draw, (x, yr(item['rate_per_min']) - 24), f"{item['rate_per_min']:.2f}", 38, bold=True, anchor='mb')
    text(draw, (x, 790), labels[i], 32, anchor='mt')
    text(draw, (x, 840), f"{item['spacing_s']:.3f} s mean spacing", 27, GRAY, anchor='mt')
text(draw, (85, 920), '66 intervals per mode. Within-run denied records and pauses remain in the elapsed spans.', 26, GRAY)
text(draw, (85, 960), 'Not a stopwatch measurement of 70 complete transactions or uninterrupted gate capacity.', 26, GRAY)
im.save(ROOT / 'Figure-Recorded-Entry-Rate.png')
print('Created two chapter-ready figures.')
