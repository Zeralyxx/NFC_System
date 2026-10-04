import hashlib
import json
from pathlib import Path
from statistics import mean, stdev

import openpyxl

folder = Path(__file__).resolve().parent
source = folder / 'Survey-Response-Entry-Template.xlsx'
workbook = openpyxl.load_workbook(source, data_only=False)
previous = json.loads((folder / 'survey-analysis.json').read_text(encoding='utf-8'))
sections = [('FS', 0, 5), ('PE', 5, 9), ('US', 9, 14), ('RE', 14, 19), ('SE', 19, 24)]
lines = [
    '# Completed Survey Entry Audit', '',
    f'Current workbook: `{source}`',
    f'SHA-256: `{hashlib.sha256(source.read_bytes()).hexdigest()}`',
    f'Earlier tally source: `{previous["source"]}`', '',
    'The original workbook was read without modifying it. Counts below follow the order **4, 3, 2, 1**.', '',
    'The current entries and earlier tallies disagree. Neither source is automatically treated as correct. Check against the original papers before replacing thesis tables or finalizing inferential results.', '',
]
for sheet_name, group_name in [('Manual', 'Manual Pre-Evaluation'), ('System', 'Students')]:
    sheet = workbook[sheet_name]
    rows = [[sheet.cell(r, c).value for c in range(2, 26)] for r in range(2, 72)]
    invalid = [(r + 2, c + 2, value) for r, row in enumerate(rows) for c, value in enumerate(row)
               if type(value) not in (int, float) or value not in (1, 2, 3, 4)]
    assert not invalid, invalid
    assert len({sheet.cell(r, 1).value for r in range(2, 72)}) == 70
    scores = [mean(mean(row[start:end]) for _, start, end in sections) for row in rows]
    checked = sum(bool(sheet.cell(r, 33).value) for r in range(2, 72))
    old = next(g for g in previous['groups'] if g['name'] == group_name)
    mismatches = []
    codes = [f'{code}{i + 1}' for code, start, end in sections for i in range(end - start)]
    for i, item in enumerate(old['items']):
        entered = [sum(row[i] == value for row in rows) for value in (4, 3, 2, 1)]
        if entered != item['frequencies']:
            mismatches.append((codes[i], openpyxl.utils.get_column_letter(i + 2), entered, item['frequencies']))
    lines += [
        f'## {sheet_name}', '',
        f'Complete forms: **70**. Valid answers: **1,680 of 1,680**. Reviewer initials recorded: **{checked} of 70**.', '',
        f'Provisional overall mean: **{mean(scores):.6f}**. Respondent-level sample SD: **{stdev(scores):.6f}**.',
        'Overall scores average the five category scores equally, using unrounded answers.', '',
        f'Questions differing from earlier tally: **{len(mismatches)} of 24**.', '',
        '| Question | Excel Column | Current Counts (4,3,2,1) | Earlier Counts (4,3,2,1) |',
        '| --- | --- | --- | --- |',
    ]
    for code, column, entered, original in mismatches:
        lines.append(f'| {code} | {column} | {", ".join(map(str, entered))} | {", ".join(map(str, original))} |')
    lines.append('')
lines += [
    '## Next Decision', '',
    'Confirm whether the current workbook was copied directly from the original individual papers and independently checked. If those papers support these entries, update the earlier tally tables to match them. If not, correct only the entries demonstrated to be wrong by the papers. Do not adjust individual responses simply to force agreement with an aggregate tally.', '',
    'After source reconciliation, check comparable survey content and scoring direction, then conduct the planned two-sided Welch independent-samples comparison of student overall scores. Keep administrator and security-personnel responses separate. Processing-time comparisons require separate assessment because existing manual and automated timing boundaries differ.', '',
]
output = folder / 'Completed-Survey-Entry-Audit.md'
output.write_text('\n'.join(lines), encoding='utf-8')
print(output)
