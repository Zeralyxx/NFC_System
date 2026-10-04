import json
import math
from pathlib import Path

import openpyxl
import pdfplumber
import pypdfium2 as pdfium
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

ROOT=Path(__file__).parent
data=json.loads((ROOT/'analysis.json').read_text())
manual=json.loads((ROOT/'manual.json').read_text())
auth_errors={'DENIED [PIN_FAILURE]','DENIED [PIN_LOCKED]','DENIED [QR_REPLACED_OR_MISMATCH]'}
auth=[r for r in data['records'] if r['category']=='Successful entry' or (r['values'][2]=='Entry' and r['values'][4] in auth_errors)]
pin=[r for r in auth if r['values'][3]!='Fast']
assert len(auth)==222 and len(pin)==152
assert sum(r['values'][4]=='DENIED [PIN_LOCKED]' for r in pin)==2
for mode,auth_n,pin_n in [('Fast',70,0),('Standard',79,79),('HighSecurity',73,73)]:
    assert sum(r['values'][3]==mode for r in auth)==auth_n
    assert sum(r['values'][3]==mode for r in pin)==pin_n
book=openpyxl.load_workbook(ROOT/'Chapter-4-Combined-Metrics.xlsx',data_only=True)
errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e']
assert not errors,errors
expected={'C56':1,'C57':152,'C58':2,'C59':2/152,'C60':18/438,'C61':manual['mean_s'],'C62':manual['service_equivalent_per_min'],'C63':1}
for cell,value in expected.items():
    assert math.isclose(book['Summary'][cell].value,value,abs_tol=1e-9),(cell,book['Summary'][cell].value,value)
for i,row in enumerate(manual['observations'],20):
    assert book['ManualBaseline'].cell(i,1).value==row['respondent']
    assert book['ManualBaseline'].cell(i,2).value==row['seconds']
for i,row in enumerate(data['records'],7):
    for j,value in enumerate(row['values'][1:],2):
        actual=book['RawRecords'].cell(i,j).value
        assert actual==value or (isinstance(value,(int,float)) and math.isclose(actual,value,abs_tol=1e-9)),(i,j,actual,value)
book.close()
pdfpath=ROOT/'Chapter-4-Results-and-Discussion.pdf'
reader=PdfReader(pdfpath)
text='\n'.join(p.extract_text() for p in reader.pages)
for forbidden in ['88.42%','96.15%','11.58%','9.38 times','2.71 times','1.66 times','revised observation','Revision status','Reserved for','latest user evaluation remains pending']:
    assert forbidden not in text,forbidden
assert all(section in text for section in ['4.1 Implementation','4.2 Results','4.3 Testing','4.4 Respondent','4.5 User','4.6 Discussion','4.7 Summary'])
assert '222' in text and '152' in text and '21.803' in text
survey=json.loads((ROOT/'chapter-survey-analysis.json').read_text())
completed=json.loads((ROOT/'completed-survey-analysis.json').read_text())
assert len([item for group in survey['groups'] for item in group['items']])==111
assert survey['individual_response_source'].endswith('Survey-Response-Entry-Template.xlsx')
for group,expected_mean in zip(survey['groups'],[3.7417142857142855,3.634,3.81375,2.2835714285714284]):
    assert math.isclose(group['overall_category_mean'],expected_mean,abs_tol=1e-10)
    for item in group['items']:
        assert item['question'] in (ROOT/'Chapter-4-Results-and-Discussion.md').read_text(encoding='utf-8')
        assert sum(item['frequencies'])==group['respondents']
for required in ['4.5.6 Comparison','Table 4.22.','Table 4.23.','Figure 4.11.','3.81','2.28','15.12','89.08','1.458','Sample SD']:
    assert required in text,required
chapter=(ROOT/'Chapter-4-Results-and-Discussion.md').read_text(encoding='utf-8')
previous=(ROOT/'Chapter-4-before-survey-update.md').read_text(encoding='utf-8')
assert chapter.split('## 4.4 Respondent Profile')[0] == previous.split('## 4.4 Respondent Profile')[0]
assert 'overall equal-category mean of 2.37' not in chapter
assert 'manual pre-evaluation mean was 2.37' not in chapter
entry=openpyxl.load_workbook(ROOT/'Survey-Response-Entry-Template.xlsx',data_only=False)
for index,name in [(0,'System'),(3,'Manual')]:
    group=survey['groups'][index]
    values=[[entry[name].cell(r,c).value for c in range(2,26)] for r in range(2,72)]
    for i,item in enumerate(group['items']):
        assert item['frequencies']==[sum(row[i]==rating for row in values) for rating in [4,3,2,1]]
    assert math.isclose(group['overall_category_mean'],completed['groups'][name]['descriptive']['Overall']['mean'],abs_tol=1e-12)
old_survey=json.loads((ROOT/'survey-analysis.json').read_text())
for index in [1,2]:
    assert [item['frequencies'] for item in survey['groups'][index]['items']] == [item['frequencies'] for item in old_survey['groups'][index]['items']]
entry.close()
out=ROOT/'full-chapter-render';out.mkdir(exist_ok=True)
doc=pdfium.PdfDocument(pdfpath)
for i in range(len(doc)):
    doc[i].render(scale=1.5).to_pil().save(out/f'page-{i+1:02}.png')
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',18)
for start in range(0,len(doc),6):
    sheet=Image.new('RGB',(1200,1160),'#E5E7E9');draw=ImageDraw.Draw(sheet)
    for j,i in enumerate(range(start,min(start+6,len(doc)))):
        image=Image.open(out/f'page-{i+1:02}.png');image.thumbnail((380,530))
        x=15+(j%3)*400;y=35+(j//3)*575
        sheet.paste(image,(x,y));draw.text((x,y-25),f'Page {i+1}',font=font,fill='black')
    sheet.save(out/f'contact-{start//6+1}.png')
with pdfplumber.open(pdfpath) as p:
    overflow=[]
    for i,page in enumerate(p.pages,1):
        for word in page.extract_words():
            if word['x0']<65 or word['x1']>547 or word['top']<20 or word['bottom']>765:
                overflow.append((i,word['text'],word['x0'],word['x1'],word['top'],word['bottom']))
    assert not overflow,overflow[:20]
print(f'Validated {len(reader.pages)} PDF pages, 70 manual records, 438 system records, 222 credential decisions and 152 PIN submissions.')
print('No workbook error cells, old performance percentages or page-boundary text overflows found.')
print('Validated all 111 survey items, group means, source selection, and inclusion in the chapter.')
print('Manual and student tallies match individual responses; security/admin counts and pre-survey chapter sections are unchanged.')
